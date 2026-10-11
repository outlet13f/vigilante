"""Build the Word deliverables from Markdown sources.

    python deliverables/build.py            # every deliverables/src/*.md
    python deliverables/build.py FILE.md    # one file

Each source starts with a front matter block:

    ---
    title: 아키텍처 설계서
    doc_id: VGL-SI-02
    version: 1.0
    date: 2026-10-11
    author: Vigilante 개발팀
    status: 초안
    history:
      - 1.0 | 2026-10-11 | 최초 작성
    ---

Supported Markdown: # to #### headings, paragraphs, **bold**, `code`,
- / * bullets (nested by two spaces), 1. numbered lists, | tables |,
``` code blocks ```, > notes, and --- page breaks. Output goes to
deliverables/docx/<source name>.docx. Needs python-docx (1.1.x).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC, OUT = ROOT / "src", ROOT / "docx"
FONT = "Malgun Gothic"
MONO = "Consolas"
ACCENT = RGBColor(0x1F, 0x3A, 0x68)


def set_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def styles(doc):
    base = doc.styles["Normal"]
    base.font.name, base.font.size = FONT, Pt(10)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    base.paragraph_format.space_after = Pt(4)
    base.paragraph_format.line_spacing = 1.25
    for level, size in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
        st = doc.styles[f"Heading {level}"]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
        st.font.color.rgb = ACCENT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def inline(par, text, size=None, mono=False):
    """Add text with **bold** and `code` spans."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            for sub in re.split(r"(`[^`]+`)", part[2:-2]):  # code inside bold
                if sub.startswith("`") and sub.endswith("`") and len(sub) > 1:
                    set_font(par.add_run(sub[1:-1]), name=MONO, size=(size or 10) - 0.5, bold=True)
                elif sub:
                    set_font(par.add_run(sub), size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            set_font(par.add_run(part[1:-1]), name=MONO, size=(size or 10) - 0.5)
        else:
            set_font(par.add_run(part), name=MONO if mono else FONT, size=size)


def shade(cell, hex_color):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def table(doc, rows):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            inline(p, text.replace("<br>", "\n"), size=9)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                shade(cells[j], "DCE3EE")
    if rows:
        tr = t.rows[0]._tr
        trpr = tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trpr.append(hdr)
    doc.add_paragraph()


def code(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, "F3F4F6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for k, line in enumerate(lines):
        r = p.add_run(line)
        set_font(r, name=MONO, size=8.5)
        if k < len(lines) - 1:
            r.add_break()
    doc.add_paragraph()


def toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
        else:
            t = OxmlElement("w:t")
            t.text = "목차는 Word에서 열 때 갱신됩니다 (F9 또는 '필드 업데이트')."
            run._r.append(t)
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


def front_matter(text):
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if not m:
        raise SystemExit("front matter (--- ... ---) missing")
    meta, history, key = {}, [], None
    for line in m.group(1).splitlines():
        if line.strip().startswith("- ") and key == "history":
            history.append([c.strip() for c in line.strip()[2:].split("|")])
        elif ":" in line:
            key, val = line.split(":", 1)
            key = key.strip()
            meta[key] = val.strip()
    meta["history"] = history
    for need in ("title", "doc_id", "version", "date"):
        if not meta.get(need):
            raise SystemExit(f"front matter needs {need}")
    return meta, text[m.end():]


def cover(doc, meta):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Vigilante — Unified Rollback Orchestrator"), size=12, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(meta["title"]), size=26, bold=True)
    for _ in range(8):
        doc.add_paragraph()
    rows = [["문서 번호", meta["doc_id"]], ["버전", meta["version"]], ["작성일", meta["date"]],
            ["작성", meta.get("author", "")], ["상태", meta.get("status", "")]]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        c = t.add_row().cells
        inline(c[0].paragraphs[0], k, size=10)
        shade(c[0], "DCE3EE")
        inline(c[1].paragraphs[0], v, size=10)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph(style="Heading 1")
    set_font(h.add_run("개정 이력"), size=16, bold=True, color=ACCENT)
    table(doc, [["버전", "일자", "내용"]] + (meta["history"] or [[meta["version"], meta["date"], "최초 작성"]]))
    h = doc.add_paragraph(style="Heading 1")
    set_font(h.add_run("목차"), size=16, bold=True, color=ACCENT)
    toc(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def header_footer(doc, meta):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(hp.add_run(f"{meta['title']}  |  {meta['doc_id']} v{meta['version']}"), size=8, color=RGBColor(0x66, 0x66, 0x66))
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.text = text
        run._r.append(el)
    set_font(run, size=8)


def render(body, doc):
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            code(doc, block)
            i += 1
            continue
        if s == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style=f"Heading {len(m.group(1))}")
            inline(p, m.group(2))
            for run in p.runs:
                run.font.color.rgb = ACCENT
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip().strip("|").replace("\\|", "\x00")  # \| is a literal pipe
                cells = [c.strip().replace("\x00", "|") for c in row.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            inline(p, s.lstrip("> "), size=9.5)
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            depth = min(len(m.group(1)) // 2, 2)
            numbered = m.group(2)[0].isdigit()
            style = ("List Number" if numbered else "List Bullet") + ("" if depth == 0 else f" {depth + 1}")
            p = doc.add_paragraph(style=style)
            inline(p, m.group(3))
            i += 1
            continue
        para = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^\s*(#{1,4}\s|[-*]\s|\d+\.\s|\||>|```)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        inline(doc.add_paragraph(), " ".join(para))


def build(path):
    meta, body = front_matter(path.read_text(encoding="utf-8").replace("\r\n", "\n"))
    doc = Document()
    styles(doc)
    header_footer(doc, meta)
    cover(doc, meta)
    render(body, doc)
    doc.core_properties.title = meta["title"]
    doc.core_properties.subject = meta["doc_id"]
    doc.core_properties.author = meta.get("author", "")
    OUT.mkdir(exist_ok=True)
    out = OUT / (path.stem + ".docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    files = [Path(a) for a in sys.argv[1:]] or sorted(SRC.glob("*.md"))
    for f in files:
        print("built", build(f))
